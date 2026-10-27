# -*- coding: utf-8 -*-
"""
方案A（主文流+富文本赋值版）：
- 仅统计主文流（wdMainTextStory）里的表格，源/目标口径一致
- 按主文流出现顺序配对：源第 n 张 ↔ 目标第 n 张
- 按“目标表在文中的起点”倒序执行，避免删除导致对象/索引漂移
- 复制逻辑采用 Word COM 富文本赋值：rng.FormattedText = src_tbl.Range.FormattedText
依赖：pip install pywin32
环境：Windows + Microsoft Word
"""

from __future__ import annotations

import contextlib
import os
import time
import shutil
from pathlib import Path
from copy import deepcopy
from typing import Dict, List, Set, Tuple

try:
    import pythoncom
except Exception:
    pythoncom = None

import win32com.client as win32
from win32com.client import constants as C

try:
    import pywintypes
except Exception:
    print("⚠️ 未找到 pywintypes，pywin32 可能安装不完整。")
    pywintypes = None



try:
    from docx import Document
    from docx.oxml.ns import qn
except Exception:
    Document = None
    qn = None



# 常量兼容处理
WD_STYLE_TYPE_TABLE = getattr(C, 'wdStyleTypeTable', 3)
WD_ORGANIZER_OBJECT_STYLES = getattr(C, 'wdOrganizerObjectStyles', 0)
WD_MAIN_TEXT_STORY = getattr(C, 'wdMainTextStory', 1)

# --- 配置 ---
TEMPLATE_DOC_PATH = "../temp/step0_preprocess_标题预处理_20251104_235356.docx"
SOURCE_DOC_PATH = "../temp/step11_library_number_高级格式化完成_20251104_235500.docx"
OUTPUT_DOC_PATH = "../output/formatted_document——1.docx"
CLEANED_DOC_PATH = "output/cleaned_document.docx"

STYLE_SYNC_ENABLED = True
STYLE_SYNC_INCLUDE_TYPES = {WD_STYLE_TYPE_TABLE}
STYLE_SYNC_EXCLUDE_PREFIXES: tuple[str, ...] = ("Heading", "TOC")
STYLE_SYNC_EXTRA_NAMES: Set[str] = set()
STYLE_SYNC_VERBOSE = False



STYLE_BACKUP_ENABLED = True

STYLE_BACKUP_TYPES = {"paragraph", "character", "table"}

# 为空则意味着所有类型都会备份（包括被自定义过的内置样式）
STYLE_BACKUP_CUSTOM_ONLY_TYPES: Set[str] = {"paragraph", "character"}
STYLE_BACKUP_LIMIT_TO_TABLE_USAGE = True
STYLE_BACKUP_INCLUDE_BASE_STYLES = False

STYLE_BACKUP_EXCLUDE_PREFIXES: tuple[str, ...] = ("Heading", "TOC")

STYLE_BACKUP_EXCLUDE: Set[str] = set()

STYLE_BACKUP_SOURCE_OVERRIDE: str | None = None

STYLE_BACKUP_VERBOSE = False

STYLE_BACKUP_EXTRA_NAMES: Set[str] = set()



RETRY_MAX = 3


def _abs(p: str) -> str:
    return str(Path(p).expanduser().resolve())


def _need_retry_com_error(e) -> bool:
    hresult = getattr(e, "hresult", None)
    if hresult in (-2147418111, -2147023174, -2147417836, -2147023170):
        return True
    msg = str(e)
    if "远程过程调用失败" in msg or "The remote procedure call failed" in msg:
        return True
    return any(k in msg for k in (
        "被呼叫方拒绝接收调用",
        "RPC 服务器不可用",
        "Call was rejected by callee",
        "The RPC server is unavailable",
    ))


def _pump_messages():
    if not pythoncom:
        return
    with contextlib.suppress(Exception):
        from pythoncom import PumpWaitingMessages

        PumpWaitingMessages()


def _retry_call(fn, *args, **kwargs):
    delay = 0.6
    for attempt in range(1, RETRY_MAX + 1):
        try:
            return fn(*args, **kwargs)
        except Exception as e:
            if pywintypes and isinstance(e, pywintypes.com_error) and _need_retry_com_error(e):
                _pump_messages()
                time.sleep(delay)
                delay = min(delay * 1.4 + 0.2, 2.5)
                continue
            raise


def _iter_collection(col):
    count = getattr(col, "Count", 0)
    for index in range(1, count + 1):
        with contextlib.suppress(Exception):
            yield col(index)


def _style_name_should_exclude(name: str) -> bool:
    if not name:
        return True
    return any(name.startswith(prefix) for prefix in STYLE_SYNC_EXCLUDE_PREFIXES)


def _copy_whitelisted_styles(app, doc_src, doc_out):
    if not STYLE_SYNC_ENABLED:
        return

    try:
        src_styles = _retry_call(lambda: doc_src.Styles)
    except Exception as err:
        print(f"⚠️ 无法读取源文档样式：{err}")
        return

    names: Set[str] = set(STYLE_SYNC_EXTRA_NAMES)
    for style in _iter_collection(src_styles):
        with contextlib.suppress(Exception):
            if getattr(style, "BuiltIn", False):
                continue
            style_type = style.Type
            if style_type not in STYLE_SYNC_INCLUDE_TYPES:
                continue
            name = str(style.NameLocal or style.Name)
            if _style_name_should_exclude(name):
                continue
            names.add(name)

    if not names:
        return

    for name in sorted(names):
        try:
            _retry_call(
                app.OrganizerCopy,
                doc_src.FullName,
                doc_out.FullName,
                name,
                WD_ORGANIZER_OBJECT_STYLES,
            )
            if STYLE_SYNC_VERBOSE:
                print(f"   · 已同步样式：{name}")
        except Exception as copy_err:
            print(f"⚠️ 样式同步失败：{name} -> {copy_err}")


def _collect_table_style_ids(doc_path: str) -> Set[str]:
    if Document is None:
        return set()
    path_obj = Path(doc_path)
    if not path_obj.exists():
        return set()
    if not _wait_for_file_access(path_obj, mode="read"):
        return set()

    doc = Document(str(path_obj))
    collected: Set[str] = set(STYLE_BACKUP_EXTRA_NAMES)

    def add_style(style):
        if style is None:
            return
        while style is not None:
            style_id = getattr(style, "style_id", None)
            if not style_id:
                break
            if any(style_id.startswith(prefix) for prefix in STYLE_BACKUP_EXCLUDE_PREFIXES):
                pass
            elif style_id not in STYLE_BACKUP_EXCLUDE:
                collected.add(style_id)
            if not STYLE_BACKUP_INCLUDE_BASE_STYLES:
                break
            style = getattr(style, "base_style", None)

    for table in doc.tables:
        add_style(getattr(table, "style", None))
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    add_style(getattr(paragraph, "style", None))
                    for run in paragraph.runs:
                        add_style(getattr(run, "style", None))

    return collected


def _backup_docx_styles(doc_path: str, allowed_ids: Set[str] | None = None) -> Dict[str, object]:
    if not STYLE_BACKUP_ENABLED:
        return {}
    if Document is None or qn is None:
        if STYLE_BACKUP_VERBOSE:
            print("⚠️ 未安装 python-docx，跳过样式备份。")
        return {}
    path_obj = Path(doc_path)
    if not path_obj.exists():
        print(f"⚠️ 样式备份失败：未找到文档 {doc_path}")
        return {}
    if not _wait_for_file_access(path_obj, mode="read"):
        print(f"⚠️ 样式备份失败：暂时无法读取 {doc_path}")
        return {}

    doc = Document(str(path_obj))
    styles_element = doc.styles.element
    collected: Dict[str, object] = {}

    for node in styles_element.findall(qn("w:style")):
        style_id = node.get(qn("w:styleId"))
        if not style_id:
            continue
        style_type = node.get(qn("w:type"))
        if allowed_ids is not None and style_id not in allowed_ids:
            continue
        if STYLE_BACKUP_TYPES and style_type not in STYLE_BACKUP_TYPES:
            continue
        if any(style_id.startswith(prefix) for prefix in STYLE_BACKUP_EXCLUDE_PREFIXES):
            continue
        if style_id in STYLE_BACKUP_EXCLUDE:
            continue
        if (
            STYLE_BACKUP_CUSTOM_ONLY_TYPES
            and style_type in STYLE_BACKUP_CUSTOM_ONLY_TYPES
            and node.get(qn("w:customStyle")) != "1"
        ):
            continue
        collected[style_id] = deepcopy(node)

    if STYLE_BACKUP_VERBOSE:
        print(f"🔒 已备份样式 {len(collected)} 项")
    return collected


def _restore_docx_styles(doc_path: str, styles: Dict[str, object]):
    if not styles:
        return
    if Document is None or qn is None:
        if STYLE_BACKUP_VERBOSE:
            print("⚠️ 未安装 python-docx，跳过样式回写。")
        return
    path_obj = Path(doc_path)
    if not path_obj.exists():
        print(f"⚠️ 样式回写失败：未找到文档 {doc_path}")
        return
    if not _wait_for_file_access(path_obj, mode="read"):
        print(f"⚠️ 样式回写失败：暂时无法读取 {doc_path}")
        return

    doc = Document(str(path_obj))
    styles_element = doc.styles.element
    existing = {
        node.get(qn("w:styleId")): node
        for node in styles_element.findall(qn("w:style"))
        if node.get(qn("w:styleId"))
    }

    for style_id, node in styles.items():
        current = existing.get(style_id)
        if current is not None:
            styles_element.remove(current)
        styles_element.append(deepcopy(node))

    if not _wait_for_file_access(path_obj, mode="write"):
        print(f"⚠️ 样式回写失败：暂时无法写入 {doc_path}")
        return
    doc.save(str(path_obj))
    if STYLE_BACKUP_VERBOSE:
        print(f"🔁 已回写样式 {len(styles)} 项")


def _get_tables_main_list(doc) -> List:
    rng = _retry_call(lambda: doc.StoryRanges(WD_MAIN_TEXT_STORY))
    tables = [t for t in _retry_call(lambda: rng.Tables)]
    tables.sort(key=lambda t: t.Range.Start)
    return tables


def _prepare_output_document(dst: str, outp: str, retries: int = 5):
    src_path = Path(dst)
    out_path = Path(outp)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    delay = 0.4
    last_error = None
    _wait_for_file_access(src_path, mode="read")

    for attempt in range(1, retries + 1):
        try:
            if out_path.exists():
                if not _wait_for_file_access(out_path, mode="write"):
                    raise PermissionError(f"timeout waiting to access {out_path}")
                with contextlib.suppress(Exception):
                    out_path.unlink()
                    _wait_for_file_access(out_path, mode="write")
            else:
                _wait_for_file_access(out_path, mode="write")
            shutil.copy2(src_path, out_path)
            _wait_for_file_access(out_path, mode="read")
            return
        except (PermissionError, OSError) as exc:
            last_error = exc
            if attempt >= retries:
                raise
            time.sleep(delay)
            delay = min(delay * 1.6, 2.0)
            continue

    if last_error:
        raise last_error


def _copy_table_direct_font(src_tbl, dst_tbl):
    if not src_tbl or not dst_tbl:
        return
    src_rows = _retry_call(lambda: src_tbl.Rows.Count)
    dst_rows = _retry_call(lambda: dst_tbl.Rows.Count)
    src_cols = _retry_call(lambda: src_tbl.Columns.Count)
    dst_cols = _retry_call(lambda: dst_tbl.Columns.Count)
    rows = min(src_rows, dst_rows)
    cols = min(src_cols, dst_cols)

    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            try:
                src_cell = _retry_call(lambda: src_tbl.Cell(r, c))
                dst_cell = _retry_call(lambda: dst_tbl.Cell(r, c))
            except Exception:
                continue
            _copy_range_font(src_cell.Range, dst_cell.Range)


def _copy_range_font(src_range, dst_range):
    try:
        dst_font = _retry_call(lambda: dst_range.Font)
        src_font = _retry_call(lambda: src_range.Font)
    except Exception:
        return

    try:
        dst_range.ParagraphFormat.Alignment = src_range.ParagraphFormat.Alignment
    except Exception:
        pass

    attributes = [
        "Name", "NameAscii", "NameFarEast", "NameOther",
        "Size", "Bold", "Italic", "Underline",
        "Color", "ColorIndex", "Spacing", "Scaling",
        "Position", "Kerning", "Shadow", "Outline",
        "Emboss", "Engrave", "StrikeThrough", "DoubleStrikeThrough",
        "Subscript", "Superscript"
    ]

    for attr in attributes:
        with contextlib.suppress(Exception):
            setattr(dst_font, attr, getattr(src_font, attr))


def _wait_for_file_access(path: Path, mode: str, timeout: float = 10.0, interval: float = 0.3) -> bool:
    deadline = time.time() + timeout
    while True:
        try:
            if mode == "read":
                if not path.exists():
                    return True
                with open(path, "rb"):
                    return True
            else:  # write
                with open(path, "a+b"):
                    return True
        except FileNotFoundError:
            if mode == "write":
                return True
        except OSError:
            pass

        if time.time() >= deadline:
            return False
        time.sleep(interval)


def _hard_replace_table(doc_out, src_tbl, dst_tbl) -> bool:
    try:
        src_fmt = _retry_call(lambda: src_tbl.Range.FormattedText)
        start_pos = _retry_call(lambda: dst_tbl.Range.Start)
        _retry_call(dst_tbl.Delete)

        rng = _retry_call(lambda: doc_out.Range(Start=start_pos, End=start_pos))
        _retry_call(lambda: setattr(rng, "FormattedText", src_fmt))

        # ====== 新增：统一表格内容为单倍行距（只动段落） ======
        try:
            new_tbl = rng.Tables(1)
            for para in new_tbl.Range.Paragraphs:
                para.LineSpacingRule = 0   # wdLineSpaceSingle
                para.LineSpacing = 12      # 可选，保险值
        except Exception:
            pass

        return True

    except Exception as e:
        print(f"    ❌ 硬替换失败：{e}")
        return False


def _create_word_app():
    try:
        app = win32.Dispatch("Word.Application")
    except Exception:
        app = win32.DispatchEx("Word.Application")

    with contextlib.suppress(Exception):
        app.Visible = False
    with contextlib.suppress(Exception):
        app.DisplayAlerts = 0
    with contextlib.suppress(Exception):
        app.AutomationSecurity = 3
        app.Options.CheckSpellingAsYouType = False
        app.Options.CheckGrammarAsYouType = False
    with contextlib.suppress(Exception):
        app.Options.SaveNormalPrompt = False
        app.Options.SavePropertiesPrompt = False
    return app


def _open_document(app, path: str, *, readonly: bool):
    documents = getattr(app, 'Documents', None)
    if documents is None:  # Word instance not ready
        raise AttributeError('Word.Application.Documents 属性不可用')
    return documents.Open(
        path,
        ReadOnly=readonly,
        AddToRecentFiles=False,
        ConfirmConversions=False,
        Revert=False,
        OpenAndRepair=False,
        NoEncodingDialog=True,
    )


def _close_doc(doc, save_changes):
    if not doc:
        return
    with contextlib.suppress(Exception):
        _retry_call(doc.Close, SaveChanges=save_changes)


def _quit_app(app):
    if not app:
        return
    with contextlib.suppress(Exception):
        _retry_call(app.Quit)

def _find_document_by_path(app, expected_path: str):
    try:
        expected = Path(expected_path).resolve()
    except Exception:
        return None
    with contextlib.suppress(Exception):
        count = app.Documents.Count
        for idx in range(1, count + 1):
            doc = app.Documents(idx)
            with contextlib.suppress(Exception):
                if Path(doc.FullName).resolve() == expected:
                    return doc
    return None



def _open_documents_with_restart(src: str, dst: str, outp: str) -> Tuple:
    """
    Open source and target documents with retries; returns (app, doc_src, doc_out).
    """
    last_error = None
    for attempt in range(1, RETRY_MAX + 1):
        app = _create_word_app()
        doc_src = None
        doc_out = None
        try:
            doc_src = _retry_call(lambda: _open_document(app, src, readonly=True))
            if doc_src is None:
                doc_src = _find_document_by_path(app, src)
            if doc_src is None:
                raise ValueError("DOC_OPEN_NONE")

            _prepare_output_document(dst, outp)
            doc_out = _retry_call(lambda: _open_document(app, outp, readonly=False))
            if doc_out is None:
                doc_out = _find_document_by_path(app, outp)
            if doc_out is None:
                raise ValueError("DOC_OPEN_NONE")

            _copy_whitelisted_styles(app, doc_src, doc_out)

            return app, doc_src, doc_out

        except Exception as e:
            last_error = e
            should_retry = (
                pywintypes and isinstance(e, pywintypes.com_error) and _need_retry_com_error(e)
            ) or (isinstance(e, ValueError) and str(e) == "DOC_OPEN_NONE")

            _close_doc(doc_out, save_changes=False)
            _close_doc(doc_src, save_changes=False)
            _quit_app(app)

            if should_retry and attempt < RETRY_MAX:
                time.sleep(0.5 * attempt)
                continue
            raise
    raise last_error


def replace_tables_in_mainstory_all(original_path: str, edited_path: str, output_path: str) -> Tuple[bool, int]:
    src = _abs(original_path)
    dst = _abs(edited_path)
    outp = _abs(output_path)

    if not Path(src).exists():
        print(f"? ԭ始文档不存在：{src}")
        return False, 0
    if not Path(dst).exists():
        print(f"? 目标文档不存在：{dst}")
        return False, 0
    Path(outp).parent.mkdir(parents=True, exist_ok=True)

    style_backup: Dict[str, object] = {}
    if STYLE_BACKUP_ENABLED:
        style_source_path = STYLE_BACKUP_SOURCE_OVERRIDE or src
        allowed_ids = None
        if STYLE_BACKUP_LIMIT_TO_TABLE_USAGE:
            allowed_ids = _collect_table_style_ids(_abs(style_source_path))
            if not allowed_ids:
                allowed_ids = None
        style_backup = _backup_docx_styles(_abs(style_source_path), allowed_ids) if (allowed_ids or STYLE_BACKUP_EXTRA_NAMES) else {}

    pythoncom_initialized = False
    if pythoncom:
        pythoncom.CoInitialize()
        pythoncom_initialized = True

    app = None
    doc_src = None
    doc_out = None
    replaced = 0
    success = False
    should_restore_styles = False

    try:
        app, doc_src, doc_out = _open_documents_with_restart(src, dst, outp)

        src_tbls = _get_tables_main_list(doc_src)
        dst_tbls = _get_tables_main_list(doc_out)
        n = min(len(src_tbls), len(dst_tbls))

        print("?? ɨ���������еı���")
        print(f"�� 源（主文流）：{len(src_tbls)} 张表")
        print(f"�� 目标（主文流）：{len(dst_tbls)} 张表")

        if n == 0:
            print("?? 主文流一侧无表格，无需替换。")
            _retry_call(doc_out.Save)
            success = True
            should_restore_styles = bool(style_backup)
        else:
            jobs = [(src_tbls[i], dst_tbls[i]) for i in range(n)]
            jobs.sort(key=lambda pair: pair[1].Range.Start, reverse=True)

            print(f"�� 计划替换：{n} 项（按目标表在文中的起点倒序执行）")
            for idx, (s, d) in enumerate(jobs[:8], 1):
                print(f"  �� #{idx} 源Start={s.Range.Start} → 目Start={d.Range.Start}")
            if n > 8:
                print(f"  �� …… 其余 {n} 项")

            for s_tbl, d_tbl in jobs:
                if _hard_replace_table(doc_out, s_tbl, d_tbl):
                    replaced += 1
                    print("     ✅ 成功")
                else:
                    print("     ❌ 失败")

            try:
                _retry_call(doc_out.Save)
            except Exception:
                _retry_call(doc_out.SaveAs, outp, FileFormat=12)

            print(f"[OK] 已保存：{outp}")
            print(f"   - 完成替换：{replaced}/{n}")
            success = True
            should_restore_styles = bool(style_backup)

    except Exception as e:
        print(f"? 执行出错：{e}")
        import traceback

        traceback.print_exc()
    finally:
        _close_doc(doc_out, save_changes=True)
        _close_doc(doc_src, save_changes=False)
        _quit_app(app)
        if pythoncom_initialized:
            with contextlib.suppress(Exception):
                pythoncom.CoUninitialize()

    if should_restore_styles and style_backup:
        try:
            _restore_docx_styles(outp, style_backup)
            if STYLE_BACKUP_VERBOSE:
                print("🔁 已回写自定义样式。")
        except Exception as style_err:
            print(f"⚠️ 样式回写失败：{style_err}")

    return success, replaced




def replace_tables(src_path: str, dst_path: str, out_path: str):
    replace_tables_in_mainstory_all(src_path, dst_path, out_path)


if __name__ == "__main__":
    replace_tables(
        TEMPLATE_DOC_PATH,
        SOURCE_DOC_PATH,
        OUTPUT_DOC_PATH,
    )
