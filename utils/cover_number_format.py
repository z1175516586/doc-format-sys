from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

# =========================
# 直接指定文件路径（自行修改）
# =========================
input_path = r"C:\Users\yanha\Desktop\原始文档.docx"
output_path = r"C:\Users\yanha\Desktop\格式化后.docx"

# 匹配模式：库号 + 冒号(中或英) + 可选空格 + 编号（字母数字）
pattern = re.compile(r'(库号[:：]\s*)([A-Za-z0-9]+)')

# 字体设置
LABEL_FONT_NAME = 'SimHei'            # 黑体
LABEL_FONT_EAST_ASIA = '黑体'
CODE_FONT_NAME = 'Times New Roman'    # 英文字体
FONT_SIZE_PT = 16                     # 三号 ≈ 16pt


def is_cjk(char):
    code = ord(char)
    return (
        0x4E00 <= code <= 0x9FFF or
        0x3400 <= code <= 0x4DBF or
        0xF900 <= code <= 0xFAFF or
        0x3000 <= code <= 0x303F or  # CJK符号与标点
        0xFF01 <= code <= 0xFF60 or  # 全角ASCII与标点（含：）
        0xFFE0 <= code <= 0xFFEF    # 全角符号
    )


def estimate_text_width(text, font_size_pt):
    width = 0.0
    for char in text:
        if char.isspace():
            width += font_size_pt * 0.5
        elif is_cjk(char):
            width += font_size_pt
        else:
            width += font_size_pt * 0.5
    return width


def update_style_value(style, key, value):
    parts = [part for part in style.split(';') if part]
    updated = []
    seen = False
    for part in parts:
        if ':' in part:
            k, v = part.split(':', 1)
            if k.strip() == key:
                updated.append(f"{key}:{value}")
                seen = True
            else:
                updated.append(part)
        else:
            updated.append(part)
    if not seen:
        updated.append(f"{key}:{value}")
    return ';'.join(updated) + ';'


def adjust_vml_line(paragraph, target_width_pt, extend_left_pt=0.0, never_shrink=True):
    v_lines = paragraph._element.findall('.//{urn:schemas-microsoft-com:vml}line')
    if not v_lines:
        return
    line = v_lines[0]
    style = line.get('style')
    if not style:
        return

    width_match = re.search(r'width:([0-9.]+)pt', style)
    margin_match = re.search(r'margin-left:([0-9.]+)pt', style)
    if not width_match or not margin_match:
        return

    old_width = float(width_match.group(1))
    old_margin_left = float(margin_match.group(1))
    right_edge = old_margin_left + old_width

    # 不缩短原线条，避免“库号”被覆盖不足
    base_width = max(old_width, target_width_pt) if never_shrink else max(target_width_pt, 1.0)
    new_width = max(base_width + float(extend_left_pt or 0.0), 1.0)
    # 固定右端，向左扩展
    new_margin_left = right_edge - new_width

    style = update_style_value(style, 'width', f"{new_width:.2f}pt")
    style = update_style_value(style, 'margin-left', f"{new_margin_left:.2f}pt")
    line.set('style', style)


def style_run(run, font_name, east_asia_name=None, size_pt=16, bold=True, underline=False):
    """设置 run 的字体样式"""
    font = run.font
    font.name = font_name
    if east_asia_name:
        run._element.rPr.rFonts.set(qn('w:eastAsia'), east_asia_name)
    font.size = Pt(size_pt)
    font.bold = bold
    font.underline = underline


def set_paragraph_bottom_border(paragraph, style="thickThinSmallGap", size=12, space=1, color="000000"):
    """为段落添加“上粗下细”双线型下边框（类似WPS/Word的边框样式）"""
    p = paragraph._element
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is None:
        p_bdr = OxmlElement('w:pBdr')
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn('w:bottom'))
    if bottom is None:
        bottom = OxmlElement('w:bottom')
        p_bdr.append(bottom)
    bottom.set(qn('w:val'), style)
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), str(space))
    bottom.set(qn('w:color'), color)


def process_paragraph(paragraph):
    """查找并格式化库号"""
    text = paragraph.text
    matches = list(pattern.finditer(text))
    if not matches:
        return

    pieces = []
    last_idx = 0
    for m in matches:
        if m.start() > last_idx:
            pieces.append(('text', text[last_idx:m.start()]))
        pieces.append(('label', m.group(1)))
        pieces.append(('code', m.group(2)))
        last_idx = m.end()
    if last_idx < len(text):
        pieces.append(('text', text[last_idx:]))

    # 清空原有文本 runs，保留图形/线条与前置缩进空白
    seen_text = False
    for r in paragraph.runs:
        r_el = r._element
        if r_el.find(qn('w:drawing')) is not None or r_el.find(qn('w:pict')) is not None:
            continue
        text = r.text or ''
        if text:
            if not seen_text and text.isspace():
                # 保留前置空白以维持模板对齐
                continue
            if text.strip():
                seen_text = True
            r.text = ''

    # 按 pieces 重新生成 runs
    for kind, content in pieces:
        run = paragraph.add_run(content)
        if kind == 'label':
            style_run(run, LABEL_FONT_NAME, LABEL_FONT_EAST_ASIA, FONT_SIZE_PT, True, underline=False)
        elif kind == 'code':
            style_run(run, CODE_FONT_NAME, None, FONT_SIZE_PT, True, underline=False)
        # 其他文本保持默认格式

    label_text = ''.join(piece[1] for piece in pieces if piece[0] == 'label')
    code_text = ''.join(piece[1] for piece in pieces if piece[0] == 'code')
    label_width = estimate_text_width(label_text, FONT_SIZE_PT)
    code_width = estimate_text_width(code_text, FONT_SIZE_PT)
    # 基线用编码宽度，避免误缩短；向左延伸覆盖“库号”
    target_width = code_width
    extend_left = label_width + (FONT_SIZE_PT * 0.3)
    adjust_vml_line(paragraph, target_width, extend_left_pt=extend_left, never_shrink=True)


def iterate_cells(doc):
    """遍历所有表格单元格中的段落"""
    for table in doc.tables:
        for row in table.rows:
            try:
                cells = row.cells
            except ValueError:
                # python-docx 在处理缺失 vMerge 起点的表格时会抛错，这里直接跳过该行
                continue
            for cell in cells:
                for p in cell.paragraphs:
                    yield p


def iterate_headers_and_footers(doc):
    """遍历页眉页脚中的段落"""
    for section in doc.sections:
        for p in section.header.paragraphs:
            yield p
        for p in section.footer.paragraphs:
            yield p


def process_document(input_path, output_path):
    """处理整个文档"""
    doc = Document(input_path)
    for p in doc.paragraphs:
        process_paragraph(p)
    for p in iterate_cells(doc):
        process_paragraph(p)
    for p in iterate_headers_and_footers(doc):
        process_paragraph(p)
    doc.save(output_path)
    print(f"格式化完成：{output_path}")


if __name__ == "__main__":
    process_document(input_path, output_path)
