# -*- coding: utf-8 -*-
"""
功能：
1. shorten_table_rows：将第一个表格所有行高设为 0.5 cm，并返回原始第二行行高 r
2. restore_table_rows：将第一个表格所有行高恢复为 r

说明：
- 文件原地保存
- r 使用 EMU 单位（python-docx 原生单位）
"""

from docx import Document
from docx.shared import Cm


def shorten_table_rows(docx_path: str) -> int:
    """
    缩短表格行高

    :param docx_path: Word 文件路径
    :return: 原第二行的行高 r（EMU）
    """
    doc = Document(docx_path)

    if not doc.tables:
        raise ValueError("文档中没有表格")

    table = doc.tables[0]

    if len(table.rows) < 2:
        raise ValueError("表格行数不足 2 行，无法获取第二行行高")

    # 获取第二行原始行高（EMU）
    second_row = table.rows[1]
    r = second_row.height

    # 如果没有显式设置行高，height 可能为 None
    if r is None:
        # 默认一个值（比如 0.8cm），避免恢复失败
        r = Cm(0.8)

    # 设置所有行高为 0.5 cm
    for row in table.rows:
        row.height = Cm(0.5)

    doc.save(docx_path)
    print(f"已缩短表格行高为 0.5cm，原始行高 r={r}")

    return r


def restore_table_rows(docx_path: str, r) -> None:
    """
    恢复表格行高

    :param docx_path: Word 文件路径
    :param r: 原始行高（EMU）
    """
    doc = Document(docx_path)

    if not doc.tables:
        raise ValueError("文档中没有表格")

    table = doc.tables[0]

    for row in table.rows:
        row.height = r

    doc.save(docx_path)
    print(f"已恢复表格行高为 r={r}")


# -------------------------
# 单独测试入口
# -------------------------
if __name__ == "__main__":
    file_path = r"D:\ASE_develop_file\doc-convertor\test\2026年4月27日.docx"  # ← 改成你的文件路径

    # Step1：缩短
    r_value = shorten_table_rows(file_path)

    input("按回车恢复表格...")

    # Step2：恢复
    restore_table_rows(file_path, r_value)