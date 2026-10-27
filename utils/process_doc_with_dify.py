import os
import shutil
import re
import json
from docx import Document
from docx.shared import Pt
import requests

def process_doc_with_dify(file_path: str) -> str:
    """
    处理文档标题：
    1. 扫描文档，收集可能的标题到列表 X
    2. 分批调用 Dify 接口识别标题
    3. 根据 Dify 返回结果设置样式
    """
    batch_size = 100
    results = []

    try:
        # 1️⃣ 复制文档到临时文件
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        temp_path = os.path.join(os.path.dirname(file_path), f"{base_name}_processed.docx")
        shutil.copy2(file_path, temp_path)
        doc = Document(temp_path)

        # 2️⃣ 扫描段落，提取可能的标题
        heading_pattern = re.compile(r'^\d+(\s*\.\s*\d+)*[\s\t]*.*$')
        X = []
        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text and heading_pattern.match(text):
                X.append({"text": text, "para_index": idx})

        if not X:
            print("⚠️ 文档中未找到可能的标题")
            return temp_path

        # Dify 配置
        DIFY_API_URL = "http://172.23.109.63:8088/v1/workflows/run"
        DIFY_API_KEY = "app-wi3L1RxqgorJrmDe5UwZrFFM"
        headers = {
            "Authorization": f"Bearer {DIFY_API_KEY}",
            "Content-Type": "application/json"
        }

        # 3️⃣ 分批调用
        for i in range(0, len(X), batch_size):
            batch = X[i:i + batch_size]

            payload = {
                "inputs": {
                    "candidates": {
                        "items": [{"text": item["text"]} for item in batch]
                    }
                },
                "response_mode": "blocking",
                "user": "test"
            }

            print(f"🚀 发送 batch {i // batch_size + 1}")
            response = requests.post(DIFY_API_URL, headers=headers, json=payload, timeout=100)

            print(response.status_code)
            print(response.text)

            if response.status_code != 200:
                print("❌ 请求失败，跳过该批次")
                continue

            result = response.json()

            raw_text = result.get("data", {}).get("outputs", {}).get("text", "")

            if not raw_text:
                print("❌ 返回为空，跳过该批次")
                continue

            clean_text = raw_text.replace("```json", "").replace("```", "").strip()

            try:
                batch_result = json.loads(clean_text)
                results.extend(batch_result)
            except Exception as e:
                print("❌ JSON解析失败，跳过该批次:", e)
                continue

        # 4️⃣ 应用结果
        level_to_style = {
            1: "Heading 1",
            2: "Heading 2",
            3: "Heading 3",
            4: "Heading 4",
            5: "Heading 5"
        }

        def normalize_number(text: str) -> str:
            """
            标准化段落开头编号：
            - 去掉所有空白字符
            - 保留数字和点
            例：'3 .4' -> '3.4'
            """
            # 只处理开头序号部分（数字+点）
            match = re.match(r'^([\d\s\.\t\u3000]+)', text)  # \u3000 是全角空格
            if match:
                num_part = match.group(1)
                # 去掉所有空白字符
                num_part_clean = re.sub(r'\s+', '', num_part)
                # 拼上剩余文本
                return num_part_clean + text[len(num_part):]
            return text

        # 应用在设置标题样式时
        for item in results:
            text_norm = normalize_number(item.get("text", ""))
            level = item.get("level", 1)

            for x in X:
                orig_text_norm = normalize_number(x["text"])
                if orig_text_norm == text_norm:
                    para = doc.paragraphs[x["para_index"]]
                    style = level_to_style.get(level, "Heading 5")
                    para.style = style
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
                    break

        # 5️⃣ 保存
        doc.save(temp_path)
        print(f"✅ 标题处理完成，文档保存至 {temp_path}")
        return temp_path

    except Exception as e:
        print(f"❌ 处理文档时出错: {str(e)}")
        return file_path


if __name__ == "__main__":
    file_path = r"D:\ASE_develop_file\2026年3月24日北航&东大程序部署调试\260331100408_鞍钢联众炼钢一线炉余热回收项目可研（20251227） - 副本.docx"

    from utils.fix_docx import fix_docx_null_relationships

    fix_docx_null_relationships(file_path)

    processed_doc = process_doc_with_dify(file_path)