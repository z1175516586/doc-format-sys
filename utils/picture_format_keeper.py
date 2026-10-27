from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

# =========================================================
# 工具函数：获取对齐
# =========================================================
def _get_align(p):
    jc = p._element.xpath('./w:pPr/w:jc')
    if not jc:
        return "left"
    return jc[0].get(qn('w:val')) or "left"

# =========================================================
# 工具函数：设置对齐
# =========================================================
def _set_align(p, align):
    mapping = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT
    }
    if align in mapping:
        p.alignment = mapping[align]

# =========================================================
# 工具函数：获取缩进
# =========================================================
def _get_indent(p):
    ind = p._element.xpath('./w:pPr/w:ind')
    if not ind:
        return 0, 0
    left = ind[0].get(qn('w:left'))
    first = ind[0].get(qn('w:firstLine'))
    return int(left or 0), int(first or 0)

# =========================================================
# 工具函数：设置缩进
# =========================================================
def _set_indent(p, left, first):
    fmt = p.paragraph_format
    if left:
        fmt.left_indent = left
    if first:
        fmt.first_line_indent = first

# =========================================================
# Step1：顺序提取图片格式
# =========================================================
def extract_and_mark_pictures(docx_path, output_path=None, start_index=1):
    doc = Document(docx_path)
    picture_info = {}
    pic_index = start_index - 1
    print(f"📄 打开文档: {docx_path}, 起始索引: {start_index}")

    for p_idx, p in enumerate(doc.paragraphs):
        picts = p._element.xpath('.//w:pict') + p._element.xpath('.//w:drawing')
        if not picts:
            continue

        for pict in picts:
            # 如果是 VML pict
            shapes = pict.xpath('.//v:shape', namespaces={'v': 'urn:schemas-microsoft-com:vml'})
            if shapes:
                for shape in shapes:
                    # 原有处理逻辑
                    ...
            else:
                # drawing 类型 (pandoc 转换后的常见结构)
                inline_shapes = pict.xpath('.//wp:inline', namespaces={
                    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
                    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'
                })
                for inline in inline_shapes:
                    cx = inline.xpath('.//a:ext/@cx',
                                      namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
                    cy = inline.xpath('.//a:ext/@cy',
                                      namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
                    width = float(cx[0]) / 12700 if cx else None  # EMU -> pt
                    height = float(cy[0]) / 12700 if cy else None
                    pic_index += 1
                    align = _get_align(p)
                    left_indent, first_indent = _get_indent(p)
                    picture_info[pic_index] = {
                        "width": width,
                        "height": height,
                        "align": align,
                        "left_indent": left_indent,
                        "first_indent": first_indent
                    }
                    print(
                        f"  ✅ 图片索引 {pic_index}, 尺寸: ({width}, {height}), 对齐: {align}, 缩进: left {left_indent}, first {first_indent}")

    if not output_path:
        output_path = docx_path
    doc.save(output_path)
    print(f"💾 保存文档: {output_path}, 总图片数量: {pic_index - start_index + 1}")
    return output_path, picture_info, pic_index

# =========================================================
# Step2：顺序恢复图片格式
# =========================================================
def restore_picture_formats(docx_path, picture_info, output_path=None):
    print("\n" + "=" * 60)
    print("🧩 开始顺序恢复图片格式")
    print(f"📄 输入文件: {docx_path}, 图片数量: {len(picture_info)}")
    print("=" * 60)

    doc = Document(docx_path)
    pic_index = 1
    total_restored = 0

    for p_idx, p in enumerate(doc.paragraphs):
        picts = p._element.xpath('.//w:pict') + p._element.xpath('.//w:drawing')
        for pict in picts:
            shapes = pict.xpath('.//v:shape', namespaces={'v': 'urn:schemas-microsoft-com:vml'})
            if shapes:
                for shape in shapes:
                    # 原有恢复逻辑
                    ...
            else:
                inline_shapes = pict.xpath('.//wp:inline', namespaces={
                    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
                    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'
                })
                for inline in inline_shapes:
                    info = picture_info.get(pic_index)
                    if not info:
                        pic_index += 1
                        continue
                    # 恢复尺寸
                    if info["width"] and info["height"]:
                        # pt -> EMU
                        cx = int(info["width"] * 12700)
                        cy = int(info["height"] * 12700)
                        ext = inline.xpath('.//a:ext',
                                           namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
                        if ext:
                            ext[0].set('cx', str(cx))
                            ext[0].set('cy', str(cy))
                    _set_align(p, info["align"])
                    _set_indent(p, info["left_indent"], info["first_indent"])
                    total_restored += 1
                    pic_index += 1

    if not output_path:
        output_path = docx_path
    doc.save(output_path)
    print(f"\n✅ 图片格式顺序恢复完成, 共恢复 {total_restored} 张图片")
    print(f"📄 输出文件: {output_path}")
    print("=" * 60)
    return True

if __name__ == "__main__":
    import os

    # 文档路径
    input_docx = r"C:\Users\YourUser\Documents\test.docx"
    temp_docx = None
    restored_docx = None

    print("🔹 Step 1：提取图片信息并标记")
    out_path, pic_info, last_index = extract_and_mark_pictures(
        docx_path=input_docx,
        output_path=temp_docx,
        start_index=1
    )

    print("\n🔹 Step 2：恢复图片格式")
    restore_picture_formats(
        docx_path=temp_docx,
        picture_info=pic_info,
        output_path=restored_docx
    )

    print("\n🎉 完成！输出文件：", restored_docx)