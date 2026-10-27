# -*- coding: utf-8 -*-
"""
=========================================================
脚本名称：table_font_setter.py

功能梗概：
本脚本用于处理 Word（.docx）文件中的表格字体设置，提供两个核心能力：

1. set_table_font_small_five(docx_path)
   - 功能：将文档中“第一个表格”的字体统一设置为“宋体，小五”
   - 适用场景：规范正文表格较小字号排版

2. set_table_font_four(docx_path)
   - 功能：将文档中“第一个表格”的字体统一设置为“宋体，四号”
   - 适用场景：标准正文表格排版

实现说明：
- 基于 python-docx 操作 Word
- 遍历表格 -> 单元格 -> 段落 -> run
- 对 run 级别设置字体（西文字体 + 中文字体）
- 中文字体通过 rFonts + eastAsia 强制设置为“宋体”

注意事项：
- 仅处理“第一个表格”
- 不会新增 run，只修改已有 run
- 若表格为空或不存在，不报错但会提示

依赖：
    pip install python-docx
=========================================================
"""

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


# =========================================================
# 工具函数：设置 run 为宋体 + 指定字号
# =========================================================
def _set_run_font(run, font_size_pt):
    """
    设置单个 run 的字体为：
    - 中文：宋体
    - 西文：宋体
    - 字号：指定 pt

    :param run: docx.text.run.Run
    :param font_size_pt: 字号（float），例如 9 表示小五
    """
    run.font.name = '宋体'
    run.font.size = Pt(font_size_pt)

    # 强制设置中文字体（关键点）
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


# =========================================================
# 函数1：设置为“宋体小五”
# =========================================================
def set_table_font_small_five(docx_path: str):
    """
    将 docx 文件中“第一个表格”的字体设置为：
        宋体 + 小五（9pt）

    :param docx_path: Word 文件绝对路径
    """
    doc = Document(docx_path)

    if not doc.tables:
        print("❌ 文档中不存在表格")
        return

    table = doc.tables[0]

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, 9)

    doc.save(docx_path)
    print("✅ 已设置为：宋体小五（第一个表格）")


# =========================================================
# 函数2：设置为“宋体四号”
# =========================================================
def set_table_font_four(docx_path: str):
    """
    将 docx 文件中“第一个表格”的字体设置为：
        宋体 + 四号（14pt）

    :param docx_path: Word 文件绝对路径
    """
    doc = Document(docx_path)

    if not doc.tables:
        print("❌ 文档中不存在表格")
        return

    table = doc.tables[0]

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, 14)

    doc.save(docx_path)
    print("✅ 已设置为：宋体四号（第一个表格）")


# =========================================================
# 主函数：独立测试入口
# =========================================================
if __name__ == "__main__":
    """
    使用说明：
    1. 修改下面的路径为你的 docx 文件
    2. 运行脚本进行测试
    """

    test_file = r"D:\ASE_develop_file\doc-convertor\app\src_files\tmpwco8aai_.docx"  # ← 改成你的文件路径

    # 测试1：设置小五
    # set_table_font_small_five(test_file)

    # 测试2：设置四号（会覆盖上一步）
    set_table_font_four(test_file)