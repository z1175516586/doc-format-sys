"""OOXML 各能力模块的单元测试。"""

from __future__ import annotations

from pathlib import Path

from lxml import etree

from app.core.namespaces import NSMAP, qn
from app.domain.ooxml.header_footer import HeaderFooterProcessor
from app.domain.ooxml.highlighting import HighlightProcessor
from app.domain.ooxml.package import DocxPackage
from app.domain.ooxml.pictures import PictureProcessor
from app.domain.ooxml.repair import DocumentRepairer
from app.domain.ooxml.sections import SectionProcessor
from app.domain.ooxml.tables import TableFont, TableProcessor
from app.domain.ooxml.toc import TocProcessor


def test_repair_adds_missing_heading_styles(sample_docx: Path, tmp_path: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        # 先删掉一个标题样式，模拟源文档缺失
        styles = pkg.xml("word/styles.xml")
        for style in list(styles.findall(qn("w:style"))):
            name = style.find(qn("w:name"))
            if name is not None and (name.get(qn("w:val")) or "").lower() == "heading 4":
                styles.remove(style)
                break

        added = DocumentRepairer(pkg).ensure_heading_styles()
        assert 4 in added

        output = tmp_path / "repaired.docx"
        pkg.save(output)

    with DocxPackage.open(output) as pkg:
        levels = DocumentRepairer(pkg)._existing_heading_levels(pkg.xml("word/styles.xml"))
        assert 4 in levels


def test_section_orientation_normalized(sample_docx: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        processor = SectionProcessor(pkg)
        processor.normalize_orientation()

        for info in processor.describe():
            if info.orientation == "landscape":
                assert info.width > info.height
            elif info.orientation == "portrait":
                assert info.height >= info.width


def test_clear_footers(sample_docx: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        cleared = HeaderFooterProcessor(pkg).clear_all_footers()
        assert cleared >= 0

        for name in pkg.iter_parts(prefix="word/footer", suffix=".xml"):
            texts = "".join(node.text or "" for node in pkg.xml(name).iter(qn("w:t")))
            assert "旧页脚内容" not in texts


def test_highlight_processor_counts(sample_docx: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        stats = HighlightProcessor(pkg).remove_all()
        assert stats["highlight"] >= 1
        assert pkg.xml("word/document.xml").xpath(".//w:highlight", namespaces=NSMAP) == []


def test_table_font_applied(sample_docx: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        processor = TableProcessor(pkg)
        count = processor.apply_font_to_first_table(TableFont.from_config("table_cover"))
        assert count > 0

        first_table = pkg.tables()[0]
        fonts = first_table.findall(f".//{qn('w:rFonts')}")
        assert fonts
        assert any(font.get(qn("w:eastAsia")) == "宋体" for font in fonts)


def test_cover_table_removal_requires_marker(tmp_path: Path) -> None:
    """缺少标记时不得删除任何表格（防止误删正文表格）。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("没有标记的文档")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "正文表格"
    path = tmp_path / "no_marker.docx"
    doc.save(path)

    with DocxPackage.open(path) as pkg:
        removed = TableProcessor(pkg).remove_tables_before_marker()
        assert removed == 0
        assert len(pkg.tables()) == 1


def test_graft_copies_toc_style_missing_from_base(sample_docx: Path, tmp_path: Path) -> None:
    from app.domain.ooxml.toc import graft_refreshed_toc

    base = tmp_path / "before.docx"
    refreshed = tmp_path / "after.docx"
    with DocxPackage.open(sample_docx) as pkg:
        pkg.add_part("word/media/image1.png", b"\x89PNG")
        TocProcessor(pkg).insert_field()
        styles = pkg.xml("word/styles.xml")
        for style in list(styles.findall(qn("w:style"))):
            if style.get(qn("w:styleId")) == "TOC3":
                styles.remove(style)
        pkg.save(base)

    with DocxPackage.open(sample_docx) as pkg:
        paragraph = etree.Element(qn("w:p"))
        ppr = etree.SubElement(paragraph, qn("w:pPr"))
        pstyle = etree.SubElement(ppr, qn("w:pStyle"))
        pstyle.set(qn("w:val"), "TOC3")
        run = etree.SubElement(paragraph, qn("w:r"))
        instr = etree.SubElement(run, qn("w:instrText"))
        instr.text = " PAGEREF _Toc1 \\h "
        pkg.body.insert(0, paragraph)
        styles = pkg.xml("word/styles.xml")
        style = etree.SubElement(styles, qn("w:style"))
        style.set(qn("w:type"), "paragraph")
        style.set(qn("w:styleId"), "TOC3")
        name = etree.SubElement(style, qn("w:name"))
        name.set(qn("w:val"), "toc 3")
        based = etree.SubElement(style, qn("w:basedOn"))
        based.set(qn("w:val"), "Normal")
        pkg.save(refreshed)

    assert graft_refreshed_toc(base, refreshed) == "grafted"

    with DocxPackage.open(refreshed) as pkg:
        ids = {style.get(qn("w:styleId")) for style in pkg.xml("word/styles.xml").findall(qn("w:style"))}
        assert "TOC3" in ids
        assert any(node.get(qn("w:val")) == "TOC3" for node in pkg.body.iter(qn("w:pStyle")))
        assert pkg.has("word/media/image1.png")


def test_graft_creates_toc_style_when_refresh_omits_it(sample_docx: Path, tmp_path: Path) -> None:
    from app.domain.ooxml.toc import graft_refreshed_toc

    base = tmp_path / "before.docx"
    refreshed = tmp_path / "after.docx"
    with DocxPackage.open(sample_docx) as pkg:
        pkg.add_part("word/media/image1.png", b"\x89PNG")
        TocProcessor(pkg).insert_field()
        styles = pkg.xml("word/styles.xml")
        for style in list(styles.findall(qn("w:style"))):
            if style.get(qn("w:styleId")) == "TOC3":
                styles.remove(style)
        pkg.save(base)
    with DocxPackage.open(sample_docx) as pkg:
        styles = pkg.xml("word/styles.xml")
        for style in list(styles.findall(qn("w:style"))):
            if style.get(qn("w:styleId")) == "TOC3":
                styles.remove(style)
        paragraph = etree.Element(qn("w:p"))
        ppr = etree.SubElement(paragraph, qn("w:pPr"))
        pstyle = etree.SubElement(ppr, qn("w:pStyle"))
        pstyle.set(qn("w:val"), "TOC3")
        run = etree.SubElement(paragraph, qn("w:r"))
        instr = etree.SubElement(run, qn("w:instrText"))
        instr.text = " PAGEREF _Toc1 \\h "
        pkg.body.insert(0, paragraph)
        pkg.save(refreshed)

    assert graft_refreshed_toc(base, refreshed) == "grafted"

    with DocxPackage.open(refreshed) as pkg:
        styles = pkg.xml("word/styles.xml")
        match = [
            style
            for style in styles.findall(qn("w:style"))
            if style.get(qn("w:styleId")) == "TOC3"
        ]
        assert len(match) == 1
        name = match[0].find(qn("w:name"))
        assert name is not None and name.get(qn("w:val")) == "toc 3"


def test_toc_insert_and_detect(sample_docx: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        processor = TocProcessor(pkg)
        assert not processor.has_toc_field()
        assert processor.insert_field()
        assert processor.has_toc_field()


def test_toc_title_replace(sample_docx: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        processor = TocProcessor(pkg)
        assert processor.replace_title("目  录")

        texts = [node.text or "" for node in pkg.body.iter(qn("w:t"))]
        assert "目  录" in texts


def test_picture_processor_no_crash_without_images(sample_docx: Path) -> None:
    with DocxPackage.open(sample_docx) as pkg:
        processor = PictureProcessor(pkg)
        assert processor.count() == 0
        assert processor.reset_indent() == 0


def test_heading_number_keeps_one_space() -> None:
    from app.domain.ooxml.numbering import format_number_label

    assert format_number_label("1 ") == "1 "
    assert format_number_label("1  ") == "1 "
    assert format_number_label("1.1 ") == "1.1 "
    assert format_number_label("1") == "1 "


def test_body_header_aligns_and_footer_line_is_full(tmp_path: Path) -> None:
    from shutil import copyfile

    from app.core.config import get_settings
    from app.services.template_service import TemplateService

    source = get_settings().resources_dir / "content_template.docx"
    target = tmp_path / "chrome.docx"
    copyfile(source, target)

    with DocxPackage.open(target) as pkg:
        header = pkg.xml("word/header1.xml")
        for ptab in list(header.iter(qn("w:ptab"))):
            tab = etree.Element(qn("w:tab"))
            ptab.getparent().replace(ptab, tab)
        TemplateService().polish_header_footer(pkg, cover_sections=0)
        header = pkg.xml("word/header1.xml")
        ptabs = list(header.iter(qn("w:ptab")))
        assert len(ptabs) == 1
        assert ptabs[0].get(qn("w:alignment")) == "right"
        assert ptabs[0].get(qn("w:relativeTo")) == "margin"
        assert list(header.iter(qn("w:tab"))) == []

        footer = pkg.xml("word/footer1.xml")
        wraps = [
            el
            for el in footer.iter()
            if el.tag.endswith("}wrap") or el.tag.endswith("}wrapSquare")
        ]
        assert wraps
        assert all((el.get("type") or "").lower() == "none" for el in wraps if el.tag.endswith("}wrap"))
        assert list(footer.iter(qn("wp:wrapSquare"))) == []

        from app.services.template_service import _release_footer_shapes

        snippet = etree.fromstring(
            '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
            '<wp:anchor behindDoc="0"><wp:wrapSquare wrapText="bothSides"/></wp:anchor></w:ftr>'
        )
        _release_footer_shapes(snippet)
        assert list(snippet.iter(qn("wp:wrapSquare"))) == []
        assert list(snippet.iter(qn("wp:wrapNone")))


def test_front_cover_header_footer_cleared(sample_docx: Path) -> None:
    from docx import Document

    from app.services.template_service import TemplateService

    document = Document(sample_docx)
    document.sections[0].header.paragraphs[0].text = "封面不该出现的页眉"
    document.save(sample_docx)

    with DocxPackage.open(sample_docx) as pkg:
        header = next(pkg.iter_parts(prefix="word/header", suffix=".xml"))
        node = next(pkg.xml(header).iter(qn("w:t")), None)
        if node is None:
            paragraph = next(pkg.xml(header).iter(qn("w:p")))
            run = etree.SubElement(paragraph, qn("w:r"))
            node = etree.SubElement(run, qn("w:t"))
        node.text = "封面不该出现的页眉"

        TemplateService().polish_header_footer(pkg, cover_sections=2)

        rels = pkg.xml("word/_rels/document.xml.rels")
        rel_map = {rel.get("Id"): rel.get("Target") for rel in rels}
        section = pkg.section_properties()[0]
        header_refs = section.findall(qn("w:headerReference"))
        footer_refs = section.findall(qn("w:footerReference"))
        assert {ref.get(qn("w:type")) for ref in header_refs} == {"even", "default", "first"}
        assert {rel_map[ref.get(qn("r:id"))] for ref in header_refs} == {"header_blank.xml"}
        assert {rel_map[ref.get(qn("r:id"))] for ref in footer_refs} == {"footer_blank.xml"}
        blank = pkg.xml("word/header_blank.xml")
        assert "".join(item.text or "" for item in blank.iter(qn("w:t"))) == ""
        bottoms = list(blank.iter(qn("w:bottom")))
        assert bottoms and bottoms[0].get(qn("w:val")) == "nil"
