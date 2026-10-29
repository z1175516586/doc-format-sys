"""源文档规范校验的单元测试。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from app.domain.ooxml.conformance import DocumentValidator
from app.domain.ooxml.cover import CoverMetadataExtractor
from app.domain.ooxml.headings import HeadingProcessor
from app.domain.ooxml.package import DocxPackage

_ISSUER = "鞍钢集团工程技术有限公司"


def _cover_doc(
    path: Path,
    page_one: list[str],
    page_two: list[str] | None = None,
    *,
    page_break: bool = True,
) -> Path:
    """封面第一页、第二页的段落，最后接设计人员名单。"""
    doc = Document()
    for text in page_one:
        doc.add_paragraph(text)
    if page_two is not None:
        if page_break:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        for text in page_two:
            doc.add_paragraph(text)
    doc.add_paragraph("各专业参加设计人员名单")
    doc.save(path)
    return path


def _type_issues(path: Path, selected: str):
    with DocxPackage.open(path) as pkg:
        return DocumentValidator(pkg).check_document_type(selected)


def _heading_doc(path: Path, items: list[tuple[str, int]]) -> Path:
    doc = Document()
    doc.add_paragraph("目      录")
    for text, level in items:
        doc.add_heading(text, level=level)
    doc.save(path)
    return path


def _heading_issues(path: Path):
    with DocxPackage.open(path) as pkg:
        report = HeadingProcessor(pkg).apply_heading_styles()
        return DocumentValidator(pkg).collect_heading_issues(report)


# ----------------------------------------------------------------------
# 封面文档类型
# ----------------------------------------------------------------------
def test_cover_type_accepts_canonical_with_suffix(tmp_path: Path) -> None:
    doc = _cover_doc(tmp_path / "cover.docx", ["某某工程项目", "初步设计说明书", _ISSUER])
    assert _type_issues(doc, "1") == []


def test_cover_type_accepts_spaced_wording(tmp_path: Path) -> None:
    doc = _cover_doc(
        tmp_path / "cover.docx",
        ["某某工程项目", "可 行 性 研 究", "鞍 钢 集 团 工 程 技 术 有 限 公 司"],
    )
    assert _type_issues(doc, "2") == []


def test_cover_issuer_missing_is_error(tmp_path: Path) -> None:
    doc = _cover_doc(tmp_path / "cover.docx", ["某某工程项目", "初步设计", "某某设计院"])
    issues = _type_issues(doc, "1")
    assert [item.code for item in issues] == ["cover_issuer_missing"]
    assert issues[0].severity == "error"


def test_cover_type_above_issuer_must_be_standard(tmp_path: Path) -> None:
    doc = _cover_doc(tmp_path / "cover.docx", ["某某工程项目", "可研设计", _ISSUER])
    issues = _type_issues(doc, "2")
    assert [item.code for item in issues] == ["front_type_invalid"]
    assert "可研设计" in issues[0].message


def test_mislabelled_front_page_without_page_break(tmp_path: Path) -> None:
    """二烧翻车机：第一页写成「可研设计」，第二页才是「可行性研究」，两页之间没有分页符。"""
    doc = _cover_doc(
        tmp_path / "cover.docx",
        ["鞍钢股份烧结厂二烧2号翻车机新增除尘器项目", "可研设计", _ISSUER, "2024年8月"],
        ["鞍钢股份烧结厂二烧2号翻车机新增除尘器项目", "可行性研究", "总设计师：吕艳平", _ISSUER],
        page_break=False,
    )
    issues = _type_issues(doc, "2")
    assert [item.code for item in issues] == ["front_type_invalid"]
    assert issues[0].severity == "error"


def test_cover_project_missing_is_error(tmp_path: Path) -> None:
    doc = _cover_doc(tmp_path / "cover.docx", ["初步设计", _ISSUER])
    issues = _type_issues(doc, "1")
    assert [item.code for item in issues] == ["front_project_missing"]


def test_cover_type_rejects_mismatch(tmp_path: Path) -> None:
    doc = _cover_doc(tmp_path / "cover.docx", ["某某工程项目", "初步设计", _ISSUER])
    issues = _type_issues(doc, "2")
    assert [item.code for item in issues] == ["document_type_mismatch"]


def test_cover_type_combines_split_secondary_line(tmp_path: Path) -> None:
    """「初步设计」与「（代可研）」分两行写，也要能识别成第 3 类。"""
    doc = _cover_doc(tmp_path / "cover.docx", ["某某工程项目", "初步设计", "（代可研）", _ISSUER])
    assert _type_issues(doc, "3") == []
    assert [item.code for item in _type_issues(doc, "1")] == ["document_type_mismatch"]


def test_page_two_type_mismatch_is_warning(tmp_path: Path) -> None:
    doc = _cover_doc(
        tmp_path / "cover.docx",
        ["某某工程项目", "初步设计", _ISSUER, "2024年8月"],
        ["某某工程项目", "可行性研究", _ISSUER],
    )
    issues = _type_issues(doc, "1")
    assert [item.code for item in issues] == ["page_two_type_mismatch"]
    assert issues[0].severity == "warning"


def test_project_name_comes_from_front_page_only(tmp_path: Path) -> None:
    """项目名只取第一页类型上方的文字：单位名单独一行，坐标残留去掉，不拼到第二页。"""
    doc = _cover_doc(
        tmp_path / "cover.docx",
        ["6032578105", "本溪玉麒麟新材料有限公司", "竖炉节能增效改造EPC总承包工程", "初步设计", _ISSUER],
        ["本溪玉麒麟新材料有限公司", "竖炉节能增效改造EPC总承包工程", "初步设计", _ISSUER],
    )
    with DocxPackage.open(doc) as pkg:
        meta = CoverMetadataExtractor(pkg).extract()
    assert meta.project_name == "本溪玉麒麟新材料有限公司\n竖炉节能增效改造EPC总承包工程"


# ----------------------------------------------------------------------
# 标题编号
# ----------------------------------------------------------------------
def test_heading_glued_number_is_warning(tmp_path: Path) -> None:
    doc = _heading_doc(tmp_path / "h.docx", [("1 概述", 1), ("1.1建设规模", 2)])
    issues = _heading_issues(doc)
    glued = [item for item in issues if item.code == "heading_number_glued"]
    assert glued and glued[0].severity == "warning"
    assert "1.1" in glued[0].message


def test_heading_jump_is_error(tmp_path: Path) -> None:
    doc = _heading_doc(tmp_path / "h.docx", [("1 概述", 1), ("1.1 规模", 2), ("1.3 内容", 2)])
    issues = _heading_issues(doc)
    errors = [item for item in issues if item.severity == "error"]
    assert [item.code for item in errors] == ["heading_jump_same_chapter"]
    assert errors[0].location


def test_heading_malformed_number_is_error(tmp_path: Path) -> None:
    doc = _heading_doc(tmp_path / "h.docx", [("1 概述", 1), ("1..1 规模", 2)])
    issues = _heading_issues(doc)
    errors = [item for item in issues if item.severity == "error"]
    assert [item.code for item in errors] == ["heading_number_invalid"]


def test_cross_chapter_jump_is_warning(tmp_path: Path) -> None:
    doc = _heading_doc(tmp_path / "h.docx", [("1 概述", 1), ("3 系统方案", 1)])
    issues = _heading_issues(doc)
    assert [item.code for item in issues] == ["heading_jump_cross_chapter"]
    assert issues[0].severity == "warning"


def test_body_list_item_not_reported(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("目      录")
    doc.add_heading("1 概述", level=1)
    doc.add_paragraph("1）受场地限制，采用烟道反应脱硫。")
    doc.add_paragraph("2）脱硫系统入口烟气参数")
    path = tmp_path / "h.docx"
    doc.save(path)
    # 正文列表项 "1）…" 不是标题，不能被误报
    assert _heading_issues(path) == []
