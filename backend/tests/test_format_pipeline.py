"""排版流水线端到端测试。

这些测试的作用是"回归基线"的最小可用版本：
只要流水线被破坏，这里必然失败。后续应扩展为 golden-file 比对。
"""

from __future__ import annotations

from pathlib import Path

from app.converters.noop import NoopConverter
from app.core.namespaces import NSMAP, qn
from app.domain.ooxml.package import DocxPackage, paragraph_text
from app.domain.ooxml.toc import TocProcessor
from app.schemas.document import ConvertOptions, FormatRequest
from app.services.format_service import PIPELINE_STEPS, DocumentFormatService


def _run(sample: Path, output: Path, **option_overrides) -> tuple[bool, Path, object]:
    # 单测用合成文档验证流水线机制，不做封面模板替换（模板替换由真实文档回归验证）
    option_overrides.setdefault("use_cover_template", False)
    options = ConvertOptions(**option_overrides)
    request = FormatRequest(
        source_path=str(sample),
        output_path=str(output),
        template_path=None,
        options=options,
    )
    service = DocumentFormatService(converter=NoopConverter())
    result = service.format(request)
    return result.success, output, result


def test_pipeline_succeeds(sample_docx: Path, tmp_path: Path) -> None:
    success, output, result = _run(sample_docx, tmp_path / "out.docx")

    assert success, f"流水线失败: {result.message}"
    assert output.exists()
    assert output.stat().st_size > 0
    assert len(result.steps) == len(PIPELINE_STEPS)
    assert all(step.ok for step in result.steps), [
        (step.name, step.message) for step in result.steps if not step.ok
    ]


def test_toc_field_inserted(sample_docx: Path, tmp_path: Path) -> None:
    success, output, _ = _run(sample_docx, tmp_path / "out.docx", insert_toc=True)
    assert success

    with DocxPackage.open(output) as pkg:
        assert TocProcessor(pkg).has_toc_field()


def test_toc_skipped_when_disabled(sample_docx: Path, tmp_path: Path) -> None:
    success, output, _ = _run(sample_docx, tmp_path / "out.docx", insert_toc=False)
    assert success

    with DocxPackage.open(output) as pkg:
        assert not TocProcessor(pkg).has_toc_field()


def test_footers_cleared(sample_docx: Path, tmp_path: Path) -> None:
    success, output, _ = _run(sample_docx, tmp_path / "out.docx")
    assert success

    with DocxPackage.open(output) as pkg:
        for name in pkg.iter_parts(prefix="word/footer", suffix=".xml"):
            text = "".join(node.text or "" for node in pkg.xml(name).iter(qn("w:t")))
            assert "旧页脚内容" not in text


def test_highlights_removed(sample_docx: Path, tmp_path: Path) -> None:
    success, output, _ = _run(sample_docx, tmp_path / "out.docx", remove_highlights=True)
    assert success

    with DocxPackage.open(output) as pkg:
        root = pkg.xml("word/document.xml")
        assert root.xpath(".//w:highlight", namespaces=NSMAP) == []


def test_cover_table_removed_content_table_kept(sample_docx: Path, tmp_path: Path) -> None:
    """标记「设计人员名单」之前的封面表格应被移除，正文表格应保留。"""
    success, output, _ = _run(sample_docx, tmp_path / "out.docx")
    assert success

    with DocxPackage.open(output) as pkg:
        tables = pkg.tables()
        assert len(tables) == 1, f"应只保留正文表格，实际 {len(tables)} 个"

        texts = [
            paragraph_text(p)
            for p in pkg.paragraphs()
        ]
        # 封面表格内容转为段落保留，不丢信息
        assert any("封面表格" in text for text in texts)


def test_section_break_created(sample_docx: Path, tmp_path: Path) -> None:
    """封面、目录、正文应形成多节结构。"""
    success, output, _ = _run(sample_docx, tmp_path / "out.docx")
    assert success

    with DocxPackage.open(output) as pkg:
        assert len(pkg.section_properties()) >= 2


def test_output_reopenable_by_python_docx(sample_docx: Path, tmp_path: Path) -> None:
    """输出必须能被 python-docx 正常打开，否则 Word 很可能报错。"""
    success, output, _ = _run(sample_docx, tmp_path / "out.docx")
    assert success

    from docx import Document

    doc = Document(str(output))
    assert len(doc.paragraphs) > 0


def test_warning_when_engine_unavailable(sample_docx: Path, tmp_path: Path) -> None:
    """无引擎时必须显式告警，不能静默认为目录页码已更新。"""
    success, _, result = _run(sample_docx, tmp_path / "out.docx")
    assert success
    assert any("目录页码" in warning for warning in result.warnings)


def test_missing_source_fails_cleanly(tmp_path: Path) -> None:
    options = ConvertOptions()
    request = FormatRequest(
        source_path=str(tmp_path / "不存在.docx"),
        output_path=str(tmp_path / "out.docx"),
        options=options,
    )
    result = DocumentFormatService(converter=NoopConverter()).format(request)

    assert not result.success
    assert result.steps[0].name == "open"
    assert not result.steps[0].ok
