"""测试配置与样例文档构造。

样例文档用 python-docx 生成，覆盖排版流水线依赖的全部结构：
封面段落、封面表格、目录标题、Heading 样式、高亮文本、页脚文本。
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

# 必须在导入 app 之前设置，否则 Settings 会指向真实数据目录
_TMP_ROOT = Path(tempfile.mkdtemp(prefix="dfs-test-"))
os.environ.setdefault("DFS_DATA_DIR", str(_TMP_ROOT / "data"))
os.environ.setdefault("DFS_LOG_DIR", str(_TMP_ROOT / "logs"))
os.environ.setdefault("DFS_TEMPLATE_DIR", str(_TMP_ROOT / "templates"))
os.environ.setdefault("DFS_LAYOUT_ENGINE", "none")

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_COLOR_INDEX  # noqa: E402


@pytest.fixture(scope="session")
def tmp_root() -> Path:
    return _TMP_ROOT


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """构造一份带完整结构的样例文档。"""
    doc = Document()

    # ---- 封面 ----
    doc.add_paragraph("某某工程项目")
    doc.add_paragraph("初步设计说明书")

    cover_table = doc.add_table(rows=3, cols=2)
    cover_table.style = "Table Grid"
    for row_idx, row in enumerate(cover_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.text = f"封面表格 {row_idx}-{col_idx}"

    doc.add_paragraph("设计人员名单")
    doc.add_paragraph("张三　李四")

    # ---- 目录标题 ----
    doc.add_paragraph("目      录")

    # ---- 正文 ----
    doc.add_heading("1 概述", level=1)
    doc.add_paragraph("本项目位于某某地区，主要建设内容如下。")

    highlighted = doc.add_paragraph()
    run = highlighted.add_run("需要清除的高亮文本")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.add_heading("1.1 建设规模", level=2)
    body_table = doc.add_table(rows=2, cols=3)
    body_table.style = "Table Grid"
    for row in body_table.rows:
        for cell in row.cells:
            cell.text = "正文表格"

    # ---- 页脚 ----
    section = doc.sections[0]
    section.footer.paragraphs[0].text = "旧页脚内容"

    target = tmp_path / "样例_初设.docx"
    doc.save(target)
    return target


@pytest.fixture
def template_docx(tmp_path: Path) -> Path:
    """构造一份带独立页眉页脚的模板。"""
    doc = Document()
    doc.add_paragraph("模板正文占位")

    section = doc.sections[0]
    section.header.paragraphs[0].text = "模板页眉：某某设计院"
    section.footer.paragraphs[0].text = "模板页脚：第 X 页"

    target = tmp_path / "模板.docx"
    doc.save(target)
    return target


@pytest.fixture
def api_client():
    """FastAPI 测试客户端。"""
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as client:
        yield client
